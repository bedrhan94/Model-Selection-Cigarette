#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Ultralytics training runs -> smoking_experiment_tracker.xlsx updater

What it does:
- Recursively scans a runs/ directory for Ultralytics training folders
- Reads results.csv and args.yaml
- Finds the best epoch using the highest mAP50-95 metric
- Pulls core metrics into the Train_Log sheet
- Updates existing rows if the same Results_CSV_Path already exists
- Appends new rows otherwise

Recommended usage:
    py -3.10 .\tools\update_experiment_tracker.py

Optional:
    py -3.10 .\tools\update_experiment_tracker.py --project-root "C:\path\to\project"
    py -3.10 .\tools\update_experiment_tracker.py --tracker ".\smoking_experiment_tracker.xlsx" --scan-root ".\runs"

Notes:
- This script is intentionally post-run based. It is more robust than trying to write live
  while training is still in progress.
- It auto-fills what it can infer. For richer experiment labels, use meaningful run names.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import math
import os
import re
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

import openpyxl
import yaml


IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".bmp", ".webp"}
TRAIN_SHEET = "Train_Log"
HEADER_ROW = 2
DATA_START_ROW = 3


def norm_path(p: Path) -> str:
    return str(p).replace("/", "\\")


def safe_float(value: Any) -> Optional[float]:
    if value is None:
        return None
    try:
        if isinstance(value, str):
            value = value.strip()
            if not value:
                return None
        out = float(value)
        if math.isnan(out):
            return None
        return out
    except Exception:
        return None


def safe_int(value: Any) -> Optional[int]:
    f = safe_float(value)
    if f is None:
        return None
    return int(round(f))


def read_yaml(path: Path) -> Dict[str, Any]:
    if not path.exists():
        return {}
    with path.open("r", encoding="utf-8") as f:
        data = yaml.safe_load(f) or {}
    return data if isinstance(data, dict) else {}


def read_results_csv(path: Path) -> List[Dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        return list(reader)


def choose_col(fieldnames: Iterable[str], target: str) -> Optional[str]:
    names = list(fieldnames)

    normalized = {name.lower().replace(" ", ""): name for name in names}

    preferred_patterns = {
        "precision": [
            "metrics/precision(b)",
            "metrics/precision",
            "precision",
        ],
        "recall": [
            "metrics/recall(b)",
            "metrics/recall",
            "recall",
        ],
        "map50_95": [
            "metrics/map50-95(b)",
            "metrics/map50-95",
            "map50-95",
            "map50_95",
        ],
        "map50": [
            "metrics/map50(b)",
            "metrics/map50",
            "map50",
        ],
        "epoch": [
            "epoch",
        ],
    }

    for pattern in preferred_patterns.get(target, []):
        key = pattern.lower().replace(" ", "")
        if key in normalized:
            return normalized[key]

    # Fallback fuzzy matching
    for name in names:
        low = name.lower()
        if target == "map50_95":
            if "map50-95" in low or "map50_95" in low:
                return name
        elif target == "map50":
            if ("map50" in low) and ("95" not in low):
                return name
        elif target == "precision":
            if "precision" in low:
                return name
        elif target == "recall":
            if low.endswith("recall") or "recall" in low:
                return name
        elif target == "epoch":
            if low.strip() == "epoch":
                return name
    return None


def best_epoch_from_results(rows: List[Dict[str, str]]) -> Dict[str, Any]:
    if not rows:
        return {}

    fieldnames = list(rows[0].keys())
    col_epoch = choose_col(fieldnames, "epoch")
    col_p = choose_col(fieldnames, "precision")
    col_r = choose_col(fieldnames, "recall")
    col_m50 = choose_col(fieldnames, "map50")
    col_m95 = choose_col(fieldnames, "map50_95")

    if not col_m95:
        raise ValueError(f"mAP50-95 column not found in results.csv columns: {fieldnames}")

    best_row = None
    best_score = None

    for row in rows:
        score = safe_float(row.get(col_m95))
        if score is None:
            continue
        if best_score is None or score > best_score:
            best_score = score
            best_row = row

    if best_row is None:
        return {}

    p = safe_float(best_row.get(col_p))
    r = safe_float(best_row.get(col_r))
    f1 = None
    if p is not None and r is not None and (p + r) > 0:
        f1 = 2 * p * r / (p + r)

    return {
        "best_epoch": safe_int(best_row.get(col_epoch)),
        "precision": p,
        "recall": r,
        "map50": safe_float(best_row.get(col_m50)),
        "map50_95": safe_float(best_row.get(col_m95)),
        "f1": f1,
    }


def resolve_path(ref: Any, base_dir: Path, project_root: Path) -> Optional[Path]:
    if ref is None:
        return None

    path = Path(str(ref))
    if path.is_absolute():
        return path

    candidates = [
        base_dir / path,
        project_root / path,
    ]
    for cand in candidates:
        if cand.exists():
            return cand.resolve()

    # Return best-effort fallback
    return (base_dir / path).resolve()


def count_images_in_path(path: Path) -> int:
    if not path.exists():
        return 0

    if path.is_dir():
        return sum(1 for p in path.rglob("*") if p.suffix.lower() in IMAGE_EXTS)

    if path.is_file():
        if path.suffix.lower() == ".txt":
            count = 0
            with path.open("r", encoding="utf-8", errors="ignore") as f:
                for line in f:
                    if line.strip():
                        count += 1
            return count
        if path.suffix.lower() in IMAGE_EXTS:
            return 1

    return 0


def count_split_images(spec: Any, base_root: Path, project_root: Path) -> Optional[int]:
    if spec is None:
        return None

    if isinstance(spec, (list, tuple)):
        total = 0
        for item in spec:
            resolved = resolve_path(item, base_root, project_root)
            if resolved is not None:
                total += count_images_in_path(resolved)
        return total

    resolved = resolve_path(spec, base_root, project_root)
    if resolved is None:
        return None
    return count_images_in_path(resolved)


def load_dataset_counts(data_arg: Any, run_dir: Path, project_root: Path) -> Tuple[Optional[int], Optional[int], Optional[int], str]:
    """
    Returns train_count, val_count, test_count, export_source
    """
    if data_arg is None:
        return None, None, None, ""

    data_path = resolve_path(data_arg, run_dir, project_root)
    export_source = norm_path(data_path) if data_path else str(data_arg)

    if not data_path or not data_path.exists():
        return None, None, None, export_source

    if data_path.is_file() and data_path.suffix.lower() in {".yaml", ".yml"}:
        data_yaml = read_yaml(data_path)
        base_root = data_path.parent
        if data_yaml.get("path"):
            base_root = resolve_path(data_yaml["path"], data_path.parent, project_root) or base_root

        train_count = count_split_images(data_yaml.get("train"), base_root, project_root)
        val_count = count_split_images(data_yaml.get("val"), base_root, project_root)
        test_count = count_split_images(data_yaml.get("test"), base_root, project_root)
        return train_count, val_count, test_count, export_source

    # Fallback: direct directory or txt passed as --data
    train_count = count_images_in_path(data_path)
    return train_count, None, None, export_source


def infer_dataset_type(text: str) -> str:
    low = text.lower()

    has_fp = any(k in low for k in ["false_positive", "falsepositive", "fp"])
    has_neg = any(k in low for k in ["hard_negative", "hardnegative", "negative", "neg"])

    if has_fp:
        return "positive+hard_negative+false_positive"
    if has_neg:
        return "positive+hard_negative"
    return "positive"


def infer_hours_or_slot(text: str) -> str:
    low = text.lower()

    for pattern in [r"\b(13-30|21-30)\b", r"\b(9-22|10-22|16-22)\b", r"\b([0-2]?\d-[0-2]?\d)\b"]:
        m = re.search(pattern, low)
        if m:
            return m.group(1)

    hours = sorted({m.group(0) for m in re.finditer(r"\b(?:9|1[0-9]|2[0-2])\b", low)}, key=lambda x: int(x))
    if hours:
        if len(hours) == 1:
            return hours[0]
        return f"{hours[0]}-{hours[-1]}"

    return ""


def build_train_command(args_data: Dict[str, Any], run_dir: Path) -> str:
    model = args_data.get("model", "")
    data = args_data.get("data", "")
    imgsz = args_data.get("imgsz", "")
    batch = args_data.get("batch", "")
    epochs = args_data.get("epochs", "")
    project = args_data.get("project", "")
    name = args_data.get("name", run_dir.name)

    parts = [
        "yolo detect train",
        f"model={model}" if model != "" else "",
        f"data={data}" if data != "" else "",
        f"imgsz={imgsz}" if imgsz != "" else "",
        f"batch={batch}" if batch != "" else "",
        f"epochs={epochs}" if epochs != "" else "",
        f"project={project}" if project != "" else "",
        f"name={name}" if name != "" else "",
    ]
    return " ".join(p for p in parts if p).strip()


def next_run_id(ws) -> str:
    max_id = 0
    for row in range(DATA_START_ROW, ws.max_row + 1):
        value = ws.cell(row, 1).value
        if isinstance(value, str):
            m = re.match(r"EXP-(\d+)$", value.strip())
            if m:
                max_id = max(max_id, int(m.group(1)))
    return f"EXP-{max_id + 1:03d}"


def find_header_map(ws) -> Dict[str, int]:
    out = {}
    for col in range(1, ws.max_column + 1):
        v = ws.cell(HEADER_ROW, col).value
        if isinstance(v, str) and v.strip():
            out[v.strip()] = col
    return out


def first_blank_row(ws) -> int:
    row = DATA_START_ROW
    while row <= ws.max_row:
        if ws.cell(row, 1).value in (None, ""):
            return row
        row += 1
    return ws.max_row + 1


def find_existing_row(ws, results_csv_path: str) -> Optional[int]:
    header = find_header_map(ws)
    col_results = header.get("Results_CSV_Path")
    if not col_results:
        return None

    target = results_csv_path.lower()
    for row in range(DATA_START_ROW, ws.max_row + 1):
        value = ws.cell(row, col_results).value
        if isinstance(value, str) and value.lower() == target:
            return row
    return None


def set_cell(ws, header_map: Dict[str, int], row: int, key: str, value: Any) -> None:
    col = header_map.get(key)
    if col is None:
        return
    ws.cell(row, col).value = value


def scan_run_dirs(scan_root: Path) -> List[Path]:
    if not scan_root.exists():
        return []

    run_dirs = []
    for results_csv in scan_root.rglob("results.csv"):
        if results_csv.is_file():
            run_dirs.append(results_csv.parent)
    return sorted(set(run_dirs))


def collect_run_info(run_dir: Path, project_root: Path) -> Dict[str, Any]:
    args_yaml = run_dir / "args.yaml"
    results_csv = run_dir / "results.csv"
    best_pt = run_dir / "weights" / "best.pt"

    args_data = read_yaml(args_yaml)
    result_rows = read_results_csv(results_csv)
    best = best_epoch_from_results(result_rows)

    train_count, val_count, test_count, export_source = load_dataset_counts(
        args_data.get("data"), run_dir, project_root
    )

    model_field = str(args_data.get("model", "")).strip()
    model_name = Path(model_field).stem if model_field else run_dir.name

    run_name = str(args_data.get("name", run_dir.name))
    descriptor = " ".join(
        [
            run_name,
            str(args_data.get("project", "")),
            str(args_data.get("data", "")),
            norm_path(run_dir),
        ]
    )

    dataset_type = infer_dataset_type(descriptor)
    hours_or_slot = infer_hours_or_slot(descriptor)

    negatives_used = "exports\\cvat_negative_merged" if "hard_negative" in dataset_type else ""
    fp_used = "exports\\cvat_false_positive_export" if "false_positive" in dataset_type else ""

    if results_csv.exists():
        stamp = dt.datetime.fromtimestamp(results_csv.stat().st_mtime).strftime("%Y-%m-%d")
    else:
        stamp = dt.datetime.now().strftime("%Y-%m-%d")

    return {
        "Train_Date": stamp,
        "Export_Source": export_source,
        "Dataset_Build": run_name,
        "Dataset_Type": dataset_type,
        "Hours_or_Slot": hours_or_slot,
        "Train_Images": train_count,
        "Val_Images": val_count,
        "Test_Images": test_count,
        "Negatives_Used": negatives_used,
        "FP_Challenge_Used": fp_used,
        "Model_Name": model_name,
        "Pretrained_Weights": model_field,
        "Img_Size": safe_int(args_data.get("imgsz")),
        "Batch": safe_int(args_data.get("batch")),
        "Epochs_Planned": safe_int(args_data.get("epochs")),
        "Best_Epoch": best.get("best_epoch"),
        "Patience": safe_int(args_data.get("patience")),
        "Optimizer": args_data.get("optimizer"),
        "LR0": safe_float(args_data.get("lr0")),
        "Device": str(args_data.get("device", "")),
        "Seed": safe_int(args_data.get("seed")),
        "Precision": best.get("precision"),
        "Recall": best.get("recall"),
        "mAP50": best.get("map50"),
        "mAP50_95": best.get("map50_95"),
        "F1": best.get("f1"),
        "Best_Weights_Path": norm_path(best_pt.resolve()) if best_pt.exists() else "",
        "Results_CSV_Path": norm_path(results_csv.resolve()),
        "Train_Command": build_train_command(args_data, run_dir),
        "Notes": "",
    }


def update_workbook(tracker_path: Path, rows_to_upsert: List[Dict[str, Any]]) -> Tuple[int, int]:
    wb = openpyxl.load_workbook(tracker_path)
    ws = wb[TRAIN_SHEET]
    header_map = find_header_map(ws)

    inserted = 0
    updated = 0

    for info in rows_to_upsert:
        existing_row = find_existing_row(ws, str(info["Results_CSV_Path"]))
        if existing_row is None:
            row = first_blank_row(ws)
            run_id = next_run_id(ws)
            set_cell(ws, header_map, row, "Run_ID", run_id)
            inserted += 1
        else:
            row = existing_row
            updated += 1

        for key, value in info.items():
            set_cell(ws, header_map, row, key, value)

    wb.save(tracker_path)
    return inserted, updated


def main() -> None:
    parser = argparse.ArgumentParser(description="Update smoking_experiment_tracker.xlsx from Ultralytics runs.")
    parser.add_argument("--project-root", default=".", help="Project root directory")
    parser.add_argument("--tracker", default=None, help="Path to smoking_experiment_tracker.xlsx")
    parser.add_argument("--scan-root", default=None, help="Directory to scan for run folders (default: <project-root>\\runs)")
    args = parser.parse_args()

    project_root = Path(args.project_root).resolve()
    tracker_path = Path(args.tracker).resolve() if args.tracker else (project_root / "smoking_experiment_tracker.xlsx")
    scan_root = Path(args.scan_root).resolve() if args.scan_root else (project_root / "runs")

    if not tracker_path.exists():
        raise FileNotFoundError(f"Tracker file not found: {tracker_path}")

    run_dirs = scan_run_dirs(scan_root)
    if not run_dirs:
        print(f"No results.csv found under: {scan_root}")
        return

    rows_to_upsert = []
    skipped = 0

    for run_dir in run_dirs:
        try:
            info = collect_run_info(run_dir, project_root)
            rows_to_upsert.append(info)
        except Exception as e:
            skipped += 1
            print(f"[SKIP] {run_dir} -> {e}")

    if not rows_to_upsert:
        print("No valid run folders were parsed.")
        return

    inserted, updated = update_workbook(tracker_path, rows_to_upsert)

    print(f"Tracker updated: {tracker_path}")
    print(f"Scanned run dirs : {len(run_dirs)}")
    print(f"Inserted rows    : {inserted}")
    print(f"Updated rows     : {updated}")
    print(f"Skipped runs     : {skipped}")


if __name__ == "__main__":
    main()
